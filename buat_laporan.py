from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"D:\ITS\5 KRIPTO\Diffie-Hellman-ElGamal-Kel-2\Laporan Simulasi Diffie-Hellman dan ElGamal.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.8); sec.bottom_margin = Inches(.75)
sec.left_margin = Inches(.9); sec.right_margin = Inches(.9)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size, color in [('Title', 24, '17365D'), ('Heading 1', 16, '17365D'), ('Heading 2', 12, '287A57')]:
    st = styles[name]; st.font.name = 'Aptos Display'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)

def p(text='', bold=False, align=None):
    para=doc.add_paragraph(); r=para.add_run(text); r.bold=bold
    if align is not None: para.alignment=align
    return para

def heading(text, level=1): doc.add_heading(text, level=level)

# Cover
for _ in range(3): doc.add_paragraph('')
para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=para.add_run('LAPORAN PROYEK KRIPTOGRAFI'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(40,122,87)
para=doc.add_paragraph(style='Title'); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
para.add_run('SIMULASI PESAN AMAN\nDIFFIE–HELLMAN DAN ELGAMAL')
p('Program untuk demonstrasi pertukaran kunci, enkripsi, dekripsi, dan pemeriksaan integritas pesan', align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(3): doc.add_paragraph('')
p('Disusun oleh:', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('1. ______________________________\n2. ______________________________\n3. ______________________________', align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(2): doc.add_paragraph('')
p('Mata Kuliah: ______________________________', align=WD_ALIGN_PARAGRAPH.CENTER)
p('Dosen Pengampu: ___________________________', align=WD_ALIGN_PARAGRAPH.CENTER)
p('Program Studi ______________________________\nInstitusi _________________________________\n2026', align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

heading('Abstrak')
p('Laporan ini membahas program simulasi pesan aman yang menggabungkan pertukaran kunci Diffie–Hellman dan enkripsi ElGamal. Program memperlihatkan pembentukan parameter publik, pertukaran nilai publik, enkripsi pesan per karakter, data yang dapat diamati penyadap, proses dekripsi oleh penerima, serta pemeriksaan tag integritas. Dari percobaan menggunakan pesan “halo”, pesan berhasil dipulihkan dan tag yang dihitung penerima sama dengan tag yang diterima. Program dibuat sebagai media pembelajaran dan bukan layanan komunikasi yang aman untuk penggunaan nyata.')
p('Kata kunci: kriptografi, Diffie–Hellman, ElGamal, enkripsi, integritas pesan.')

heading('1. Pendahuluan')
heading('1.1 Latar Belakang',2)
p('Pesan yang dikirim melalui jaringan dapat diamati oleh pihak lain. Kriptografi membantu menjaga kerahasiaan pesan dan mendeteksi perubahan data. Diffie–Hellman memungkinkan dua pihak memperoleh rahasia bersama melalui pertukaran nilai publik, sedangkan ElGamal menggunakan pasangan kunci publik dan privat untuk mengenkripsi serta mendekripsi pesan.')
p('Proyek ini membuat simulasi berbasis web agar tahapan tersebut dapat diamati secara langsung, termasuk data ciphertext yang terlihat oleh pihak ketiga bernama Oscar.')
heading('1.2 Rumusan Masalah',2)
p('Bagaimana menerapkan simulasi pertukaran kunci Diffie–Hellman dan enkripsi ElGamal, serta menunjukkan proses dekripsi dan pemeriksaan integritas pesan melalui antarmuka web?')
heading('1.3 Tujuan',2)
p('Tujuan proyek adalah memahami alur kerja dasar Diffie–Hellman dan ElGamal, menampilkan informasi setiap tahapan perhitungan, serta mengamati perbedaan antara pesan asli dan ciphertext yang terlihat di jaringan.')
heading('1.4 Batasan',2)
p('Program berjalan secara lokal, menggunakan parameter demo 32-bit, dan memakai satu kunci penerima demo. Nama pengirim dan penerima dapat diubah sebagai label tampilan. Sistem ditujukan untuk pembelajaran, bukan untuk mengamankan komunikasi nyata.')

heading('2. Dasar Teori')
heading('2.1 Diffie–Hellman',2)
p('Diffie–Hellman adalah metode untuk menghasilkan rahasia bersama dengan bertukar nilai publik. Dengan parameter prima p dan generator g, pihak pertama memilih rahasia a dan mengirim A = g^a mod p. Pihak kedua memilih rahasia b dan mengirim B = g^b mod p. Keduanya menghitung nilai yang sama, yaitu K = B^a mod p = A^b mod p = g^(ab) mod p.')
heading('2.2 ElGamal',2)
p('Pada ElGamal, penerima memilih kunci privat x dan menghitung kunci publik y = g^x mod p. Untuk setiap nilai pesan m, pengirim memilih bilangan acak k, lalu menghitung c1 = g^k mod p dan c2 = m × y^k mod p. Penerima menghitung s = c1^x mod p dan memulihkan m dengan m = c2 × s^(-1) mod p.')
heading('2.3 Integritas Pesan',2)
p('Program membandingkan tag yang dikirim dengan tag yang dihitung ulang dari pesan hasil dekripsi. Jika keduanya sama, tampilan menandai pesan sebagai utuh. Implementasi tag pada program ini menggunakan checksum sederhana yang dikalikan dengan kunci sesi; ini hanya demonstrasi dan bukan MAC kriptografis yang kuat.')

heading('3. Perancangan dan Implementasi')
heading('3.1 Alur Program',2)
p('Program membangkitkan parameter publik p dan g, lalu melakukan pertukaran nilai publik Diffie–Hellman untuk menghasilkan kunci sesi. Penerima membuat pasangan kunci ElGamal. Pengirim mengenkripsi setiap karakter dan membuat tag. Antarmuka menampilkan ciphertext kepada Oscar, kemudian penerima mendekripsi pesan dan membandingkan tag.')
heading('3.2 Perangkat Lunak',2)
p('Backend dibuat dengan Python menggunakan server HTTP dan JSON bawaan. Antarmuka menggunakan HTML, CSS, dan JavaScript. Program tidak memerlukan framework atau pustaka kriptografi eksternal.')
heading('3.3 Berkas Program',2)
for s in ['crypto_core.py: perhitungan matematika modular, pembangkitan parameter, Diffie–Hellman, dan ElGamal.', 'server.py: server lokal dan layanan API untuk menyiapkan sesi serta memproses pesan.', 'index.html dan style.css: struktur dan tampilan antarmuka.', 'app.js: interaksi formulir, pengiriman pesan, dan penampilan rincian tahapan.']:
    doc.add_paragraph(s, style='List Bullet')
heading('3.4 Cara Menjalankan',2)
p('Jalankan server dengan perintah python server.py dari folder proyek, kemudian buka http://localhost:8000 pada browser. Isi nama pengirim, penerima, dan pesan, lalu tekan tombol “Kirim pesan”. Rincian perhitungan dapat dibuka melalui bagian “Lihat rincian tahapan Diffie–Hellman & ElGamal”.')

heading('4. Hasil dan Pembahasan')
p('Pengujian dilakukan dengan mengirim pesan “halo” dari Tiara kepada Diva. Tampilan menunjukkan pesan berhasil diproses, hasil dekripsi “halo”, dan status “Terverifikasi”. Pada contoh tangkapan layar, parameter yang tampil adalah p = 2468727623, q = 1234363811, dan g = 5. Nilai rahasia sesi yang dihitung Tiara dan Diva sama, yaitu K = 2401847657.')
p('Pada panel Oscar, data yang ditampilkan berupa pasangan ciphertext (c1, c2) untuk tiap karakter dan tag integritas. Setelah tombol “Coba baca sebagai Oscar” digunakan, program menjelaskan bahwa Oscar tidak memiliki kunci privat penerima. Rincian tahap dekripsi menampilkan nilai pemulihan karakter h, a, l, dan o. Tag yang diterima dan tag hasil perhitungan Diva sama-sama bernilai 299057311 pada contoh tersebut.')
heading('4.1 Ringkasan Pengujian',2)
t=doc.add_table(rows=1, cols=3); t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['Skenario','Hasil pada program','Kesimpulan']): c.text=txt
for row in [
    ('Mengirim pesan “halo”','Pesan tampil kembali setelah dekripsi','Enkripsi dan dekripsi demo berjalan'),
    ('Melihat panel Oscar','Ciphertext dan tag ditampilkan','Isi asli tidak ditampilkan pada panel penyadap'),
    ('Memeriksa tahap Diffie–Hellman','Nilai K kedua pihak sama','Pertukaran kunci berhasil pada simulasi'),
    ('Memeriksa tag','Tag diterima sama dengan tag hitungan Diva','Pesan ditandai utuh dan terverifikasi')]:
    cells=t.add_row().cells
    for c,txt in zip(cells,row): c.text=txt

heading('5. Kesimpulan dan Saran')
p('Program berhasil mendemonstrasikan alur dasar pertukaran kunci Diffie–Hellman, enkripsi dan dekripsi ElGamal, serta pemeriksaan tag integritas. Pada contoh pesan “halo”, penerima memperoleh kembali pesan yang sama dan tag cocok. Panel Oscar membantu memperlihatkan bahwa data yang dipertukarkan berbeda dari teks asli.')
p('Untuk pengembangan, program dapat menggunakan pustaka kriptografi standar, pembangkit bilangan acak kriptografis, ukuran kunci yang memadai, dan mekanisme autentikasi pesan yang telah teruji. Parameter demo dan checksum saat ini tidak boleh digunakan untuk komunikasi nyata.')

heading('Daftar Pustaka')
p('Stallings, W. (2017). Cryptography and Network Security: Principles and Practice (7th ed.). Pearson.')
p('Dokumentasi dan kode sumber proyek Simulasi Pesan Aman: Diffie–Hellman & ElGamal (file README.md dan crypto_core.py).')

heading('Lampiran Dokumentasi Program')
p('Masukkan tangkapan layar yang disediakan pada bagian ini saat finalisasi laporan. Gunakan keterangan berikut agar hasil percobaan mudah dibaca:')
for s in [
    'Gambar 1. Tampilan pesan yang berhasil didekripsi dan berstatus terverifikasi.',
    'Gambar 2. Panel data yang dapat diamati Oscar dan hasil percobaan membaca ciphertext.',
    'Gambar 3. Parameter publik dan hasil pertukaran kunci Diffie–Hellman.',
    'Gambar 4. Kunci ElGamal penerima dan tabel proses enkripsi.',
    'Gambar 5. Ciphertext yang terlihat di jaringan.',
    'Gambar 6. Proses dekripsi pesan oleh penerima.',
    'Gambar 7. Hasil pemeriksaan tag integritas pesan.'
]:
    doc.add_paragraph(s, style='List Bullet')

# Footer with page number field
for section in doc.sections:
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=fp.add_run('Laporan Simulasi Kriptografi  •  '); run.font.size=Pt(8); run.font.color.rgb=RGBColor(110,120,130)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

doc.save(OUT)
print(OUT)
